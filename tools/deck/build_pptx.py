"""Session 2 as a native PowerPoint deck, builds included.

Every slide is laid out in Chromium with the fonts PowerPoint will use
(Arial for text, Courier New for the numerals), then each box, text block and
image becomes a native shape at the same place. A build step in the web deck
becomes one click in PowerPoint: a Fade entrance, the shapes of that step
together. Text stays editable; animations show in the Animation Pane.
"""
import json, sys, pathlib, copy
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from playwright.sync_api import sync_playwright
from PIL import Image

ROOT = pathlib.Path(sys.argv[1])            # folder holding project/deck.json
OUT = pathlib.Path(sys.argv[2])
ASSETS = {}
class _Any(dict):
    def __missing__(self, k): return sys.argv[3]
ASSETS = _Any()
EMU = 6350                                   # 1920 px == 13.333 in
PT = 0.5                                     # 1 px == 0.5 pt on this canvas
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
SHELL = """<!doctype html><html><head><meta charset="utf-8"><style>
html,body{margin:0;padding:0}
section{position:relative;width:1920px;height:1080px;box-sizing:border-box;overflow:hidden}
section *{box-sizing:border-box}
aside{display:none} h1,h2,h3,p{margin:0}
*{font-family:Arial,'Liberation Sans',sans-serif !important}
*[style*="JetBrains"], *[style*="JetBrains"] *{font-family:'Courier New','Liberation Mono',monospace !important}
</style></head><body>__S__</body></html>"""
EXTRACT = pathlib.Path("/tmp/claude-0/extract.js").read_text()

def e(px): return Emu(int(round(px * EMU)))
def rgb(h): return RGBColor.from_string(h)

deck = json.loads((ROOT / "project/deck.json").read_text())
prs = Presentation()
prs.slide_width, prs.slide_height = e(1920), e(1080)
blank = prs.slide_layouts[6]
report = []

def _plain(sh):
    """No theme style: every fill and line is stated on the shape itself, so
    no renderer adds a shadow or an outline the web deck does not have."""
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)
    sh.shadow.inherit = False


def add_box(slide, it, name):
    x, y, w, h = it["x"], it["y"], it["w"], it["h"]
    sides = it["sides"]
    uniform = all(abs(s["w"] - sides[0]["w"]) < .01 and s["c"] == sides[0]["c"] for s in sides)
    shapes = []
    if it["fill"] or (uniform and sides[0]["w"] > 0):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if it["radius"] > 0 else MSO_SHAPE.RECTANGLE
        sh = slide.shapes.add_shape(kind, e(x), e(y), e(w), e(h))
        if it["radius"] > 0:
            sh.adjustments[0] = min(0.5, it["radius"] / max(1, min(w, h)))
        if it["fill"]:
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(it["fill"])
        else:
            sh.fill.background()
        if uniform and sides[0]["w"] > 0 and sides[0]["c"]:
            sh.line.color.rgb = rgb(sides[0]["c"]); sh.line.width = Pt(sides[0]["w"] * PT)
        else:
            sh.line.fill.background()
        _plain(sh)
        sh.name = name; shapes.append(sh)
    if not uniform:
        # a single painted edge - the left rule on a card, a row's hairline
        for i, s in enumerate(sides):
            if s["w"] <= 0 or not s["c"]: continue
            geo = [(x, y, w, s["w"]), (x + w - s["w"], y, s["w"], h),
                   (x, y + h - s["w"], w, s["w"]), (x, y, s["w"], h)][i]
            r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *map(e, geo))
            r.fill.solid(); r.fill.fore_color.rgb = rgb(s["c"]); r.line.fill.background()
            _plain(r)
            r.name = name; shapes.append(r)
    return shapes

def add_text(slide, it, name):
    x, y, w, h = it["x"], it["y"], it["w"], it["h"]
    pad = it["pad"]
    one_line = it["lines"] <= 1 and not any(r.get("br") for r in it["runs"])
    # Chromium measured this with Liberation Sans, which has Arial's metrics,
    # so the browser's line breaks are PowerPoint's: no extra width needed.
    slack = 0
    if it["align"] == "right": x -= slack
    elif it["align"] == "center": x -= slack / 2
    tb = slide.shapes.add_textbox(e(x), e(y), e(w + slack), e(h))
    tf = tb.text_frame
    tf.word_wrap = not one_line
    tf.auto_size = None
    tf.margin_top, tf.margin_right, tf.margin_bottom, tf.margin_left = (
        e(pad[0]), e(pad[1]), e(pad[2]), e(pad[3]))
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = {"right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}.get(it["align"], PP_ALIGN.LEFT)
    p.line_spacing = Pt(it["lh"] * PT)
    p.space_before = Pt(0); p.space_after = Pt(0)
    for r in it["runs"]:
        if r.get("br"):
            p = tf.add_paragraph(); p.alignment = tf.paragraphs[0].alignment
            p.line_spacing = Pt(it["lh"] * PT); p.space_before = Pt(0); p.space_after = Pt(0)
            continue
        run = p.add_run(); run.text = r["t"]
        f = run.font
        f.size = Pt(round(r["size"] * PT * 2) / 2)
        f.bold = r["bold"]
        f.name = "Courier New" if r["mono"] else "Arial"
        if r["color"]: f.color.rgb = rgb(r["color"])
        if r["ls"]:
            run._r.get_or_add_rPr().set("spc", str(int(round(r["ls"] * PT * 100))))
    tb.name = name
    return [tb]

def add_img(slide, it, name):
    path = ASSETS[it["src"]]
    iw, ih = Image.open(path).size
    scale = min(it["w"] / iw, it["h"] / ih)
    w, h = iw * scale, ih * scale
    x, y = it["x"] + (it["w"] - w) / 2, it["y"] + (it["h"] - h) / 2
    pic = slide.shapes.add_picture(path, e(x), e(y), e(w), e(h))
    pic.name = name
    return [pic]

P = "http://schemas.openxmlformats.org/presentationml/2006/main"

def timing_xml(steps):
    """steps: [[spid, ...], ...] in click order -> <p:timing> element."""
    ids = iter(range(1, 10_000))
    nxt = lambda: str(next(ids))
    root_id, seq_id = nxt(), nxt()
    clicks = []
    for spids in steps:
        effects = []
        for i, spid in enumerate(spids):
            node = "clickEffect" if i == 0 else "withEffect"
            effects.append(f'''<p:par><p:cTn id="{nxt()}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:set><p:cBhvr><p:cTn id="{nxt()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{nxt()}" dur="500"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>
</p:childTnLst></p:cTn></p:par>''')
        clicks.append(f'''<p:par><p:cTn id="{nxt()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{nxt()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{''.join(effects)}</p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn></p:par>''')
    xml = f'''<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="{root_id}" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="{seq_id}" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{''.join(clicks)}</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
    return etree.fromstring(xml)

def transition_xml(kind):
    child = '<p:push dir="l"/>' if kind == "push" else '<p:fade/>'
    return etree.fromstring(f'<p:transition xmlns:p="{P}" spd="med">{child}</p:transition>')

with sync_playwright() as pw:
    browser = pw.chromium.launch(executable_path=CHROME)
    page = browser.new_page(viewport={"width": 1920, "height": 1080})
    for sid in deck["order"]:
        html = (ROOT / f"project/slides/{sid}.html").read_text()
        page.set_content(SHELL.replace("__S__", html), wait_until="load")
        data = page.evaluate(EXTRACT)
        slide = prs.slides.add_slide(blank)
        if data["bg"]:
            slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(data["bg"])
        groups = {}
        for n, it in enumerate(data["items"]):
            name = f"b{it['build']}-{it['kind']}-{n}"
            made = {"box": add_box, "text": add_text, "img": add_img}[it["kind"]](slide, it, name)
            if it["build"]:
                groups.setdefault(it["build"], []).extend(s.shape_id for s in made)
        if data["notes"]:
            slide.notes_slide.notes_text_frame.text = data["notes"]
        sld = slide._element
        ext = sld.find(qn("p:extLst"))
        anchor = sld.find(qn("p:clrMapOvr"))
        insert_at = list(sld).index(anchor) + 1 if anchor is not None else len(sld)
        if data["transition"]:
            sld.insert(insert_at, transition_xml(data["transition"])); insert_at += 1
        if groups:
            steps = [groups[k] for k in sorted(groups)]
            sld.insert(insert_at, timing_xml(steps))
        report.append((sid, len(data["items"]), [len(groups[k]) for k in sorted(groups)]))
    browser.close()

prs.save(OUT)
for sid, n, g in report:
    print(f"{sid:<16} {n:>3} items   clicks {len(g):>2}  shapes per click {g}")
