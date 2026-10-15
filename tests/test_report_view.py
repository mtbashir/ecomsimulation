"""What the instructor lets teams see: report sections month by month, every
month's report, and rankings - on the console from month 1, to teams once
published."""
import io
import sqlite3
import sys
sys.path.insert(0, "src")

import pytest

from ecomsim import bootstrap, params as P, scoring
from ecomsim.engine import run_round
from ecomsim.web import db, service
from ecomsim.web.app import create_app
from ecomsim.web.titlecase import tc

HEADINGS = {"products": "Sales by product this month", "buyers": "Who bought each product",
            "stock": "Stock by product, end of month", "campaigns": "Paid media, campaign by campaign"}


@pytest.fixture
def game(tmp_path):
    """Three teams, two months run, month 3 open."""
    path = tmp_path / "g.db"
    pw = db.init(path, "G", 3, rounds=12, admin_password="admin-pw", start_mode="going_concern")
    con = db.connect(path)
    for m in (1, 2):
        db.set_game(con, open_round=m)
        service.process_round(con)
    db.set_game(con, open_round=3)
    app = create_app(path)
    app.config["TESTING"] = True
    return app, pw, con


def _client(app, who, pw):
    c = app.test_client()
    c.post("/login", data={"username": who, "password": pw})
    if who != "admin":
        c.post("/brief")
    return c


def _hide(con, month, *keys):
    shown = db.report_view(con, month)
    db.set_report_view(con, month, {k: (v and k not in keys) for k, v in shown.items()})


# --- Report sections ------------------------------------------------------------------

def test_every_section_is_shown_by_default_and_rankings_from_month_four(game):
    _, _, con = game
    for m in (1, 3):
        view = db.report_view(con, m)
        assert all(view[k] for k in HEADINGS) and view["rankings"] is False
    assert db.report_view(con, 4)["rankings"] is True
    assert db.hidden_sections(con, 1) == set()


@pytest.mark.parametrize("key", list(HEADINGS))
def test_a_hidden_section_leaves_the_team_report(game, key):
    app, pw, con = game
    _hide(con, 2, key)
    team = _client(app, "team_01", pw["team_01"]).get("/results/2").data.decode()
    assert tc(HEADINGS[key]) not in team
    for other in set(HEADINGS) - {key}:
        assert tc(HEADINGS[other]) in team
    # month 1 is untouched
    assert tc(HEADINGS[key]) in _client(app, "team_01", pw["team_01"]).get("/results/1").data.decode()
    # the instructor still sees it, labelled
    admin = _client(app, "admin", "admin-pw").get("/admin/report/team_01/2").data.decode()
    assert tc(HEADINGS[key]) in admin and "hidden from the team this month" in admin


def test_hidden_buyers_leave_the_sales_summary_too(game):
    app, pw, con = game
    team = _client(app, "team_01", pw["team_01"])
    assert "Your orders came from" in team.get("/results/2").data.decode()
    _hide(con, 2, "buyers")
    assert "Your orders came from" not in team.get("/results/2").data.decode()


def test_hidden_stock_and_sales_leave_the_decision_form(game):
    _, _, con = game
    rows = service.purchase_rows(con, "team_01", {})
    assert rows["has_stock"] and any(r["close"] is not None for r in rows["rows"])
    shelf = service.monthly_catalogue(con, "team_01", {})
    assert any(r["last"] for r in shelf)
    _hide(con, 2, "stock", "products")
    rows = service.purchase_rows(con, "team_01", {})
    assert not rows["has_stock"] and all(r["close"] is None for r in rows["rows"])
    assert not any(r["last"] for r in service.monthly_catalogue(con, "team_01", {}))


def test_hidden_campaigns_leave_the_campaigns_page(game):
    _, _, con = game
    assert service.campaign_desk(con, "team_01")["last"]
    _hide(con, 2, "campaigns")
    assert service.campaign_desk(con, "team_01")["last"] == []


def test_the_team_workbook_leaves_out_what_the_report_hid(game):
    openpyxl = pytest.importorskip("openpyxl")
    _, _, con = game
    _hide(con, 2, "products", "stock")
    data, _ = service.team_workbook(con, "team_01", 2, for_team=True)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "Products" not in wb.sheetnames and "Stock" not in wb.sheetnames
    months = {r[0].value for r in wb["Products by month"].iter_rows(min_row=5) if r[0].value}
    assert months == {1}, "month 2's rows are left out"
    full, _ = service.team_workbook(con, "team_01", 2)
    assert "Products" in openpyxl.load_workbook(io.BytesIO(full)).sheetnames


def test_the_admin_grid_saves_each_month(game):
    app, _, con = game
    admin = _client(app, "admin", "admin-pw")
    page = admin.get("/admin").data.decode()
    assert "What Teams See in Their Monthly Report" in page
    ticks = [f"{k}:{m}" for m in range(1, 13) for k, _ in db.SECTIONS
             if not (m == 3 and k in ("stock", "buyers"))
             and (k != "rankings" or m >= 4)]
    admin.post("/admin/report-view", data={"show": ticks, "months": ",".join(map(str, range(1, 13)))})
    assert db.hidden_sections(con, 3) == {"stock", "buyers"}
    assert db.hidden_sections(con, 2) == set() and db.hidden_sections(con, 4) == set()


def test_a_past_game_saved_before_the_settings_shows_everything(tmp_path):
    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.row_factory = sqlite3.Row
    con.execute("CREATE TABLE game (id INTEGER)")
    assert all(db.report_view(con, 2)[k] for k in HEADINGS)
    assert db.rankings_published(con, 3) == 0


def test_a_new_game_starts_with_everything_shown(game, tmp_path):
    _, _, con = game
    _hide(con, 1, "stock")
    db.new_game(con, tmp_path / "g.db", "Next", 3, 12, "founding", keep_overrides=False)
    assert db.hidden_sections(con, 1) == set()


# --- Every month's report -------------------------------------------------------------

def test_a_team_opens_every_month_and_each_links_the_others(game):
    app, pw, _ = game
    team = _client(app, "team_01", pw["team_01"])
    for m in (1, 2):
        body = team.get(f"/results/{m}").data.decode()
        assert f"Month {m} &middot;" in body
        other = 3 - m
        assert f'href="/results/{other}"' in body
    assert team.get("/results/3").status_code == 404
    home = team.get("/").data.decode()
    assert 'href="/results/1"' in home and 'href="/results/2"' in home


# --- Rankings -------------------------------------------------------------------------

def test_the_console_ranks_from_month_one(game):
    app, _, _ = game
    admin = _client(app, "admin", "admin-pw")
    body = admin.get("/admin/console/1").data.decode()
    assert "Leaderboard &middot; Cumulative" in body and "Withheld" not in body
    assert "Teams cannot see rankings yet" in body
    assert "Month Rank" in body


def test_teams_see_rankings_only_once_published(game):
    app, pw, con = game
    team = _client(app, "team_01", pw["team_01"])
    assert team.get("/rankings").status_code == 404
    assert 'href="/rankings"' not in team.get("/").data.decode()
    shown = db.report_view(con, 2)
    db.set_report_view(con, 2, {**shown, "rankings": True})
    body = team.get("/rankings").data.decode()
    assert "Month 2 on Its Own" in body and "Cumulative Score" in body
    assert 'href="/rankings?month=1"' in body, "every month up to the published one"
    assert "Month 1 on Its Own" in team.get("/rankings?month=1").data.decode()
    assert 'href="/rankings"' in team.get("/").data.decode()
    console = _client(app, "admin", "admin-pw").get("/admin/console/2").data.decode()
    assert "Teams can see rankings up to month 2" in console


def test_rankings_hold_both_scores_and_their_ranks(game):
    _, _, con = game
    months = service.rankings(con)
    assert [m["month"] for m in months] == [1, 2]
    rows = months[-1]["rows"]
    assert [r["rank"] for r in rows] == sorted(r["rank"] for r in rows)
    assert rows[0]["rank"] == 1 and min(r["month_rank"] for r in rows) == 1
    assert all(r["move"] is not None for r in rows)
    assert all(r["move"] is None for r in months[0]["rows"])
    world = db.load_world(con, 2)
    params = service.load_params(con)
    for r in rows:
        team = world.teams[r["team_id"]]
        assert r["score"] == pytest.approx(scoring.final_score(team, params, world.teams)["total"])


def test_the_month_score_reads_that_month_alone():
    params = P.load({"n_teams": 3, "events_enabled": 0})
    world = bootstrap.new_world(params, run_id="ms")
    ids = list(world.teams)
    run_round(world, params, {})
    run_round(world, params, {ids[0]: {"2.2": 0.15}})
    run_round(world, params, {})
    one = {t: scoring.month_score(world.teams[t], params, world.teams, 1)["total"] for t in ids}
    assert max(one.values()) - min(one.values()) < 1e-6, "identical teams, identical month 1"
    two = {t: scoring.month_score(world.teams[t], params, world.teams, 2)["total"] for t in ids}
    assert two[ids[0]] < two[ids[1]], "the month a team gave 15% away scores lower"
    # reading month 2 does not depend on what came after it
    again = bootstrap.new_world(params, run_id="ms")
    run_round(again, params, {})
    run_round(again, params, {ids[0]: {"2.2": 0.15}})
    assert scoring.month_score(again.teams[ids[0]], params, again.teams, 2)["total"] == \
        pytest.approx(two[ids[0]])


def test_the_console_shows_results_round_by_round(game):
    app, _, con = game
    admin = _client(app, "admin", "admin-pw")
    body = admin.get("/admin/console/2").data.decode()
    assert "Results by Round" in body
    assert "<th>R1</th><th>R2</th>" in body
    for label in ("Cumulative rank", "Each round on its own", "Net revenue",
                  "Contribution after marketing", "Market share"):
        assert tc(label) in body
    assert 'href="/admin/console/1"' in body, "every round's console is a click away"
    # an earlier round's console stops at that round
    first = admin.get("/admin/console/1").data.decode()
    assert "<th>R1</th>" in first and "<th>R2</th>" not in first
    assert 'href="/admin/console/1"' in admin.get("/admin").data.decode()


def test_equal_scores_share_a_rank():
    assert scoring.ranks({"a": 60.0, "b": 60.0, "c": 58.0}) == {"a": 1, "b": 1, "c": 3}
    assert scoring.ranks({"a": 50.0, "b": 51.0}) == {"a": 2, "b": 1}


# --- The Excel workbook, month by month ------------------------------------------

def test_the_workbook_is_closed_to_teams_until_ticked(game):
    app, pw, con = game
    assert not db.workbook_open(con, 2)
    team = _client(app, "team_01", pw["team_01"])
    assert ".xlsx" not in team.get("/results/2").data.decode()
    assert team.get("/results/2.xlsx").status_code == 404


def test_ticking_a_month_opens_that_months_workbook_only(game):
    app, pw, con = game
    db.set_report_view(con, 2, {**db.report_view(con, 2), "workbook": True})
    team = _client(app, "team_01", pw["team_01"])
    assert "/results/2.xlsx" in team.get("/results/2").data.decode()
    assert "/results/2.xlsx" in team.get("/").data.decode(), "and on the dashboard"
    r = team.get("/results/2.xlsx")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    assert team.get("/results/1.xlsx").status_code == 404, "month 1 was not ticked"


def test_a_hidden_section_stays_out_of_the_teams_workbook(game):
    import io
    import openpyxl
    app, pw, con = game
    for m in (1, 2):
        db.set_report_view(con, m, {**db.report_view(con, m), "stock": False,
                                    "workbook": True})
    team = _client(app, "team_01", pw["team_01"])
    book = openpyxl.load_workbook(io.BytesIO(team.get("/results/2.xlsx").data))
    assert "Stock" not in book.sheetnames
    admin = _client(app, "admin", "admin-pw")
    full = openpyxl.load_workbook(io.BytesIO(admin.get("/admin/report/team_01/2.xlsx").data))
    assert "Stock" in full.sheetnames, "the instructor's copy keeps everything"


def test_the_workbook_row_is_on_the_admin_grid_and_saves(game):
    app, _, con = game
    admin = _client(app, "admin", "admin-pw")
    assert "Excel workbook to download" in admin.get("/admin").data.decode()
    months = ",".join(str(m) for m in range(1, 13))
    ticks = [f"{k}:{m}" for m in range(1, 13) for k, _ in db.SECTIONS
             if db.report_view(con, m)[k]] + ["workbook:2"]
    admin.post("/admin/report-view", data={"months": months, "show": ticks})
    assert db.workbook_open(con, 2) and not db.workbook_open(con, 3)
    assert db.hidden_sections(con, 3) == set(), "an unticked workbook hides no section"
